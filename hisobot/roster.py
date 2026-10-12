"""Station roster import from .xlsx (spec section 7, «Импорт»): all rows or none, errors by row number."""

import zipfile
from io import BytesIO

from django.core.exceptions import ValidationError
from django.db import transaction
from openpyxl import load_workbook

from .models import Mtu, Station, validate_station_name

MAX_BYTES = 1024 * 1024
MAX_UNZIPPED = 20 * 1024 * 1024  # a real roster of 1000 rows unpacks to well under 1 MB
MAX_ROWS = 1000
EXCEL_ROWS = 1_048_576  # Excel's own last row: a file cannot claim rows beyond it
HEADER = ["мту", "станция номи"]


class RosterError(Exception):
    def __init__(self, errors):
        self.errors = errors
        super().__init__("\n".join(errors))


def _text(value):
    # Any run of whitespace, a no-break space or a line break typed inside a cell included, becomes one space.
    return "" if value is None else " ".join(str(value).split())


def _read(file):
    """The header and the non-blank rows after it as (row number, MTU, name), cells trimmed."""
    data = file.read(MAX_BYTES + 1)
    if len(data) > MAX_BYTES:
        raise RosterError(["Файл ҳажми 1 МБ дан ошмаслиги керак."])
    try:
        # The upload limit counts compressed bytes; a few hundred KB can unpack to gigabytes.
        with zipfile.ZipFile(BytesIO(data)) as archive:
            if sum(info.file_size for info in archive.infolist()) > MAX_UNZIPPED:
                raise ValueError("unpacks too large")
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
        try:
            sheet = workbook.worksheets[0]
            sheet.reset_dimensions()  # read-only mode trusts the file's <dimension>, which some exporters write stale
            # max_row: a forged row number would otherwise make openpyxl yield every blank row up to it.
            cells = sheet.iter_rows(max_col=2, max_row=EXCEL_ROWS, values_only=True)
            header = [_text(value).lower() for value in next(cells, ())]
            rows = []
            for number, row in enumerate(cells, start=2):
                mtu, name = (_text(value) for value in (*row, None, None)[:2])
                if mtu or name:
                    rows.append((number, mtu, name))
                    if len(rows) > MAX_ROWS:
                        break
        finally:
            workbook.close()
    except Exception:  # noqa: BLE001 - a foreign or broken file fails in zipfile, the XML parser or openpyxl
        raise RosterError(["Файлни ўқиб бўлмади. Excel (.xlsx) файлини юкланг."]) from None
    return header, rows


def import_roster(file, *, user) -> int:
    """Add the (MTU, name) pairs that are new; returns how many stations were added."""
    header, rows = _read(file)
    if header != HEADER:
        raise RosterError(["1-қатор: сарлавҳа «МТУ | Станция номи» бўлиши керак."])
    if len(rows) > MAX_ROWS:
        raise RosterError([f"Файлда {MAX_ROWS} тадан ортиқ станция бўлмаслиги керак."])

    mtus = {mtu.name.strip().lower(): mtu for mtu in Mtu.objects.all()}
    errors, wanted, seen = [], [], {}
    for number, mtu_name, name in rows:
        mtu = mtus.get(mtu_name.lower())
        if not mtu_name:
            errors.append(f"{number}-қатор: МТУ кўрсатилмаган")
        elif mtu is None:
            errors.append(f"{number}-қатор: «{mtu_name}» номли МТУ топилмади")
        elif not user.is_boshqarma and mtu.pk != user.mtu_id:
            errors.append(f"{number}-қатор: фақат «{user.mtu.name}» МТУ станцияларини қўшиш мумкин")
        else:
            try:
                validate_station_name(name)
            except ValidationError as error:
                errors.append(f"{number}-қатор: {error.messages[0]}")
                continue
            key = (mtu.pk, name.lower())
            if key in seen:
                errors.append(f"{number}-қатор: «{name}» станцияси {seen[key]}-қаторда ҳам бор")
                continue
            seen[key] = number
            wanted.append((mtu, name))
    if errors:
        raise RosterError(errors)

    with transaction.atomic():
        mtu_ids = sorted({mtu.pk for mtu, _ in wanted})
        # Imports and single additions to one MTU take turns, so names and sort numbers see a settled roster.
        list(Mtu.objects.select_for_update().filter(pk__in=mtu_ids).order_by("pk"))
        existing, top = set(), dict.fromkeys(mtu_ids, 0)
        for mtu_id, name, sort in Station.objects.filter(mtu_id__in=mtu_ids).values_list("mtu_id", "name", "sort"):
            existing.add((mtu_id, name.lower()))
            top[mtu_id] = max(top[mtu_id], sort)
        added = []
        for mtu, name in wanted:
            if (mtu.pk, name.lower()) not in existing:
                top[mtu.pk] += 1
                added.append(Station(mtu=mtu, name=name, sort=top[mtu.pk]))
        Station.objects.bulk_create(added)
    return len(added)
