"""The daily table in the customer's Excel template (spec section 10, «Eksport»).

Values only, never formulas; station names are always text cells. A period nobody entered is a blank cell.
"""

from copy import copy
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill

from . import periods, reports

TEMPLATE = Path(__file__).resolve().parent / "xlsx" / "shablon.xlsx"
CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
FIRST_ROW = 4  # rows 1-3: the title and the two header rows
COLUMNS = 9
GREY = "FF808080"  # carried remaining
RED = "FFC00000"  # negative remaining
YELLOW = "FFFFE9A3"  # remaining typed by hand; --override-bg of app.css


def _total(value, partial):
    """The daily total: a number, `12*` as text when an expected period is missing, blank when nothing was entered."""
    if value is None:
        return None
    return f"{value}*" if partial else value


class _Sheet:
    def __init__(self, sheet, ended):
        self.sheet = sheet
        self.ended = ended  # (night, day): has the period ended? A running one has blank totals, as on the page.
        header = [sheet.cell(3, c) if sheet.cell(3, c).value is not None else sheet.cell(2, c) for c in range(1, 10)]
        # The template's header rows give the alignment; its prepared (empty) data rows give the border and font.
        sample = [sheet.cell(FIRST_ROW + 1, c) for c in range(1, 10)]
        self.alignment = [copy(cell.alignment) for cell in header]
        self.border = [copy(cell.border) for cell in sample]
        self.font = [copy(cell.font) for cell in sample]
        sheet.delete_rows(FIRST_ROW, sheet.max_row)
        self.row = FIRST_ROW

    def put(self, col, value, *, bold=False, colour=None, text=False):
        cell = self.sheet.cell(self.row, col)
        if text:
            # Control characters cannot be stored in xlsx; a name saved with one must not break the whole file.
            value = ILLEGAL_CHARACTERS_RE.sub("", value)
        cell.value = value
        if text:
            cell.data_type = "s"  # a name like "=1+1" would otherwise be written as a formula
        cell.alignment = copy(self.alignment[col - 1])
        cell.border = copy(self.border[col - 1])
        font = copy(self.font[col - 1])
        font.b = bold
        if colour:
            font.color = colour
        cell.font = font
        return cell

    def remaining(self, col, cell, *, bold=False):
        value = cell.remaining
        colour = RED if value is not None and value < 0 else GREY if cell.carried else None
        written = self.put(col, value, bold=bold, colour=colour)
        if cell.entry and cell.entry.overridden:
            written.fill = PatternFill("solid", fgColor=YELLOW)
            written.comment = Comment(f"Ҳисобланган: {cell.entry.remaining_calc}", "Дон ҳисоботи")

    def title(self, text):
        self.put(1, text, bold=True, text=True)
        for col in range(2, COLUMNS + 1):
            self.put(col, None)
        self.sheet.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=COLUMNS)
        self.row += 1

    def station(self, number, row):
        self.put(1, number)
        self.put(2, row.station.name, text=True)
        for col, cell in ((3, row.night), (6, row.day)):
            self.put(col, cell.entry.unloaded if cell.entry else None)
            self.put(col + 1, cell.entry.arrived if cell.entry else None)
            self.remaining(col + 2, cell)
        self.put(9, _total(row.total_unloaded, row.partial))
        self.row += 1

    def total(self, label, totals):
        self.put(1, None, bold=True)
        self.put(2, label, bold=True, text=True)
        numbers = (
            totals.n_unloaded,
            totals.n_arrived,
            totals.n_remaining,
            totals.d_unloaded,
            totals.d_arrived,
            totals.d_remaining,
        )
        for col, value in enumerate(numbers, start=3):
            if not self.ended[(col - 3) // 3]:
                value = None
            self.put(col, value, bold=True, colour=RED if value is not None and value < 0 else None)
        self.put(9, _total(totals.total_unloaded, totals.partial), bold=True)
        self.row += 1

    def note(self, text):
        cell = self.sheet.cell(self.row, 1, text)
        cell.data_type = "s"
        cell.font = copy(self.font[0])
        self.row += 1


def export_daily(stations, report_date, *, user, now) -> bytes:
    every = reports.daily_rows(stations, report_date, now=now)
    rows = reports.shown(every)  # as on the daily page: no rows of dashes for switched-off stations
    groups = reports.by_mtu(rows)
    totals = reports.totals(rows)

    book = load_workbook(TEMPLATE)
    sheet = book.active
    sheet["A1"] = f"{sheet['A1'].value} {report_date:%d.%m.%Y}"
    out = _Sheet(sheet, [pe <= now for pe in periods.ends_for(report_date)])
    if len(groups) > 1:
        for mtu, group, group_totals in groups:
            out.title(f"{mtu.name} МТУ")
            for number, row in enumerate(group, start=1):
                out.station(number, row)
            out.total(f"{mtu.name} МТУ жами", group_totals)
    else:
        for number, row in enumerate(rows, start=1):
            out.station(number, row)
    out.total("Жами", totals)

    out.row += 1
    out.note(f"Яратилган: {periods.local(now):%d.%m.%Y %H:%M}")
    out.note(f"Аккаунт: {user.display_name}")
    out.note(f"Киритмаган станциялар сони: {totals.missing}")
    if len(every) > len(rows):
        out.note(f"Ўчирилган {len(every) - len(rows)} та станция кўрсатилмади")
    out.note("12* — бир давр киритилмаган, сутка жами тўлиқ эмас")
    out.note(
        "Кулранг — кўчирилган қолдиқ, сариқ — қўлда ўзгартирилган (изоҳда ҳисобланган қиймат), қизил — манфий қолдиқ"
    )
    # The template has no print layout: its nine columns spill over two portrait pages.
    sheet.print_title_rows = "1:3"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True

    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()
