"""Task 10: an MTU account never sees, reaches or changes another MTU's stations (spec section 8).

REGISTRY lists every URL under /xodim/; a staff URL missing from it fails the first test.
MTU B's objects carry numbers that appear nowhere else, so any trace of them in a response is a leak:
station «Бета-бекат» (id 987654), MTU «Бетаобод» (id 424242), a hand-typed remaining of 7777.
"""

import re
from datetime import timedelta
from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import URLPattern, reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from hisobot import access, balance, periods, urls
from hisobot.models import Change, Entry, Mtu, Station, User

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
LEAKS = ["Бета-бекат", "Бетаобод", "987654", "424242", "7777"]
PIN = re.compile(r"\d{4} \d{4}")  # a printed PIN is random digits: it may contain a leak marker by chance


def roster_file(mtu_name, station_name):
    workbook = Workbook()
    workbook.active.append(("МТУ", "Станция номи"))
    workbook.active.append((mtu_name, station_name))
    buffer = BytesIO()
    workbook.save(buffer)
    return SimpleUploadedFile("stansiyalar.xlsx", buffer.getvalue(), content_type=XLSX)


def text_of(response):
    """What a person could read in the response: the page, or every cell and comment of an Excel file."""
    if response.get("Content-Type", "").startswith(XLSX):
        workbook = load_workbook(BytesIO(response.content))
        return "\n".join(
            f"{cell.value} {cell.comment.text if cell.comment else ''}"
            for sheet in workbook for row in sheet.iter_rows() for cell in row
        )
    return PIN.sub("", response.content.decode())


# name -> requests that exercise it with one side's objects: (method, url, data) built from a Side.
REGISTRY = {
    "staff_now": lambda s: [("get", reverse("staff_now"), {"pe": s.pe})],
    "staff_daily": lambda s: [("get", reverse("staff_daily"), {"d": s.day})],
    "staff_entry": lambda s: [
        ("get", s.entry_url, {}),
        ("post", s.entry_url, {"unloaded": 2, "arrived": 2, "override": "", "version": s.version, "prev": s.prev}),
    ],
    "staff_report": lambda s: [("get", reverse("staff_report"), {"from": s.week_ago, "to": s.day})],
    "staff_export": lambda s: [("get", reverse("staff_export"), {"d": s.day})],
    "stations": lambda s: [
        ("get", reverse("stations"), {}),
        ("post", reverse("stations"), {"mtu": s.mtu.pk, "name": "Янги-бекат"}),
    ],
    "station_edit": lambda s: [
        ("get", s.url("station_edit"), {}),
        ("post", s.url("station_edit"), {"form": "station", "name": s.station.name, "sort": 5}),
        ("post", s.url("station_edit"), {"form": "opening", "opening": 9}),
    ],
    # Twice: off, then on again, so the station is active for the calls after it.
    "station_toggle": lambda s: [("post", s.url("station_toggle"), {})] * 2,
    "station_pin": lambda s: [("post", s.url("station_pin"), {})],
    "station_unlock": lambda s: [("post", s.url("station_unlock"), {})],
    "stations_import": lambda s: [
        ("get", reverse("stations_import"), {}),
        ("post", reverse("stations_import"), {"file": roster_file(s.mtu.name, "Импорт-бекат")}),
    ],
    "pin_sheet": lambda s: [("post", reverse("pin_sheet"), {"station": [s.station.pk]})],
    "password_change": lambda s: [
        ("get", reverse("password_change"), {}),
        ("post", reverse("password_change"), {"current": "нотўғри", "new": "x", "again": "y"}),
    ],
    "mtu_passwords": lambda s: [
        ("get", reverse("mtu_passwords"), {}),
        ("post", reverse("mtu_passwords"), {"user": s.user.pk}),
    ],
}

# What MTU A gets when it aims a request at MTU B's objects.
FOREIGN_STATUS = {
    "staff_entry": 404,
    "station_edit": 404,
    "station_toggle": 404,
    "station_pin": 404,
    "station_unlock": 404,
    "pin_sheet": 404,
    "mtu_passwords": 404,  # boshqarma only
}


class Side:
    """One MTU's objects, as the registry needs them."""

    def __init__(self, mtu, user, station, pe):
        self.mtu, self.user, self.station = mtu, user, station
        self.pe = periods.to_slug(pe)
        day = periods.label(pe)[0]
        self.day, self.week_ago = day.isoformat(), (day - timedelta(days=6)).isoformat()
        self.entry_url = reverse("staff_entry", args=[station.pk, self.pe])
        entry = station.entries.get(period_end=pe)
        self.version, self.prev = entry.version, balance.prev_for(station, pe)[0]

    def url(self, name):
        return reverse(name, args=[self.station.pk])


@override_settings(PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher", "hisobot.access.PinHasher"])
class IsolationTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        now = timezone.now()
        cls.pe = periods.last_ended(now)
        earlier = periods.prev_end(cls.pe)
        boss = User.objects.create_user("boshqarma")
        cls.mtu_a = Mtu.objects.create(name="Альфаобод", sort=1)
        cls.mtu_b = Mtu.objects.create(pk=424242, name="Бетаобод", sort=2)
        cls.user_a = User.objects.create_user("alfa", password="alfa-parol-1", mtu=cls.mtu_a)
        cls.user_b = User.objects.create_user("beta", password="beta-parol-1", mtu=cls.mtu_b)
        start = now - timedelta(days=20)
        cls.station_a = Station.objects.create(mtu=cls.mtu_a, name="Альфа-бекат", sort=1, active_from=start)
        cls.station_b = Station.objects.create(pk=987654, mtu=cls.mtu_b, name="Бета-бекат", sort=1, active_from=start)

        # Per station: an entry, an overridden entry after it (the «Эътибор талаб» list), then a correction
        # of the first one that cascades into the second (a history row with a cause), a PIN and a lock.
        for station, (opening, unloaded, arrived, override) in [
            (cls.station_a, (10, 1, 2, 5)),
            (cls.station_b, (5000, 311, 313, 7777)),
        ]:
            write = dict(station_id=station.pk, arrived=arrived, actor=boss, ip="10.0.0.1", now=now)
            balance.save_entry(**write, period_end=earlier, unloaded=unloaded, opening=opening,
                               form_version=0, form_prev=None)
            station.refresh_from_db()
            balance.save_entry(**write, period_end=cls.pe, unloaded=unloaded, override=override,
                               form_version=0, form_prev=balance.prev_for(station, cls.pe)[0])
            balance.save_entry(**write, period_end=earlier, unloaded=unloaded + 1, form_version=1, form_prev=opening)
            access.set_pin(station, actor=boss, ip="10.0.0.1")
            Station.objects.filter(pk=station.pk).update(locked_until=now + timedelta(hours=1))

    def setUp(self):
        self.a = Side(self.mtu_a, self.user_a, self.station_a, self.pe)
        self.b = Side(self.mtu_b, self.user_b, self.station_b, self.pe)

    def login(self, user):
        self.client.force_login(user)
        session = self.client.session
        session["staff_login_at"] = timezone.now().timestamp()
        session.save()

    def call(self, method, url, data, **extra):
        if method == "get":
            return self.client.get(url, {**data, "mtu": self.mtu_b.pk}, **extra)  # ?mtu=<B> is ignored for A
        return self.client.post(url, data, **extra)

    def assertNoLeak(self, response):
        text = text_of(response)
        for marker in LEAKS:
            self.assertNotIn(marker, text)

    def state_of(self, mtu=None):
        stations = Station.objects.filter(mtu=mtu) if mtu else Station.objects.all()
        return (
            list(stations.order_by("pk").values()),
            list(Entry.objects.filter(station__in=stations).order_by("pk").values()),
            Change.objects.filter(station__in=stations).count(),
            list(User.objects.filter(mtu=mtu).values_list("password", flat=True)),
        )

    def test_registry_covers_every_staff_url(self):
        staff = {p.name for p in urls.urlpatterns if isinstance(p, URLPattern) and str(p.pattern).startswith("xodim/")}
        self.assertEqual(set(REGISTRY), staff)

    def test_own_objects_show_nothing_of_the_other_mtu(self):
        self.login(self.user_a)
        for name, requests in REGISTRY.items():
            for method, url, data in requests(self.a):
                with self.subTest(name=name, method=method, data=data):
                    response = self.call(method, url, data, follow=True)
                    self.assertEqual(response.status_code, 404 if name == "mtu_passwords" else 200)
                    self.assertNoLeak(response)

    def test_the_other_mtus_objects_are_404_and_change_nothing(self):
        before = self.state_of(self.mtu_b)
        self.login(self.user_a)
        for name, requests in REGISTRY.items():
            for method, url, data in requests(self.b):
                with self.subTest(name=name, method=method, data=data):
                    response = self.call(method, url, data)
                    if name in FOREIGN_STATUS:
                        self.assertEqual(response.status_code, FOREIGN_STATUS[name])
                    self.assertNoLeak(response)
        self.assertEqual(self.state_of(self.mtu_b), before)
        # The add form and the import name MTU B; an MTU account adds to its own MTU or nothing.
        self.assertTrue(Station.objects.filter(mtu=self.mtu_a, name="Янги-бекат").exists())
        self.assertFalse(Station.objects.filter(name="Импорт-бекат").exists())

    def test_anonymous_and_station_sessions_are_sent_to_staff_login(self):
        station_session = {"station_id": self.station_a.pk, "pin_version": self.station_a.pin_version}
        before = self.state_of()
        for who in ["anonymous", "station"]:
            self.client.logout()
            if who == "station":
                session = self.client.session
                session.update(station_session)
                session.save()
            for name, requests in REGISTRY.items():
                for method, url, data in requests(self.a):
                    with self.subTest(who=who, name=name, method=method):
                        response = self.call(method, url, data)
                        self.assertRedirects(response, reverse("staff_login"), fetch_redirect_response=False)
        self.assertEqual(self.state_of(), before)

    def test_a_staff_session_is_no_station_session(self):
        self.login(self.user_a)
        for url in [reverse("station_home"), reverse("station_entry", args=[self.a.pe])]:
            with self.subTest(url=url):
                self.assertRedirects(self.client.get(url), reverse("home"), fetch_redirect_response=False)
                response = self.client.post(url, {"unloaded": 0, "arrived": 0, "version": 1, "prev": 0})
                self.assertRedirects(response, reverse("home"), fetch_redirect_response=False)
