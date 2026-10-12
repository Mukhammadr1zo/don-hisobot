from datetime import date, datetime, timedelta
from io import BytesIO

import openpyxl
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from hisobot import access, export, periods, reports, summary
from hisobot.models import Entry, Mtu, Station, User
from hisobot.periods import TZ
from hisobot.summary import SummaryRow

D_FROM, D_TO = date(2026, 9, 10), date(2026, 9, 12)
LATER = datetime(2026, 10, 1, 8, tzinfo=TZ)  # every period of the fixtures ended long ago


def tk(day, hour):
    """Tashkent wall-clock time in September 2026."""
    return datetime(2026, 9, day, hour, tzinfo=TZ)


class Base(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.mtu_a = Mtu.objects.create(name="Тошкент", sort=1)
        cls.mtu_b = Mtu.objects.create(name="Қўқон", sort=2)

    @classmethod
    def station(cls, name, mtu, sort=0, **fields):
        fields.setdefault("active_from", tk(1, 20))
        return Station.objects.create(mtu=mtu, name=name, sort=sort, **fields)

    @staticmethod
    def entry(station, period_end, unloaded=0, arrived=0, calc=0, override=None):
        return Entry.objects.create(
            station=station,
            period_end=period_end,
            unloaded=unloaded,
            arrived=arrived,
            remaining_calc=calc,
            remaining_override=override,
        )

    def scope(self):
        return Station.objects.order_by("mtu__sort", "sort", "id")


class SummaryTests(Base):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.a1 = cls.station("Олмазор", cls.mtu_a, 1, opening_remaining=10, opening_at=tk(9, 8), active_from=tk(9, 8))
        cls.entry(cls.a1, tk(9, 20), unloaded=1, arrived=2, calc=11)  # before the range: stock only
        cls.entry(cls.a1, tk(10, 8), unloaded=3, arrived=1, calc=9)
        cls.entry(cls.a1, tk(11, 20), unloaded=2, arrived=5, calc=12, override=10)
        cls.entry(cls.a1, tk(12, 20), unloaded=4, arrived=0, calc=6)
        cls.entry(cls.a1, tk(13, 8), unloaded=9, arrived=9, calc=6)  # after the range
        cls.a2 = cls.station("Сергели", cls.mtu_a, 2, active_from=tk(11, 9))  # expected from 11.09 20:00 on
        cls.b1 = cls.station("Бета", cls.mtu_b, 1)
        cls.entry(cls.b1, tk(12, 8), unloaded=5, arrived=7, calc=-2)

    def test_summary_sums_the_range_and_counts_periods_per_station_and_per_mtu(self):
        rows, per_mtu = summary.summary(self.scope(), D_FROM, D_TO, now=LATER)
        self.assertEqual(
            rows,
            [
                SummaryRow(self.a1, self.mtu_a, 9, 6, 6, 3, 6, 1),
                SummaryRow(self.a2, self.mtu_a, None, None, None, 0, 3, 0),  # nothing entered: not a zero
                SummaryRow(self.b1, self.mtu_b, 5, 7, -2, 1, 6, 0),
            ],
        )
        self.assertEqual(
            per_mtu,
            [SummaryRow(None, self.mtu_a, 9, 6, 6, 3, 9, 1), SummaryRow(None, self.mtu_b, 5, 7, -2, 1, 6, 0)],
        )
        self.assertEqual([r.missing for r in rows], [3, 3, 5])
        self.assertEqual(summary.total(rows), SummaryRow(None, None, 14, 13, 4, 4, 15, 1))

    def test_a_period_not_ended_at_now_is_not_expected(self):
        rows, per_mtu = summary.summary(Station.objects.filter(pk=self.a2.pk), D_FROM, D_TO, now=tk(12, 10))
        self.assertEqual(rows, [SummaryRow(self.a2, self.mtu_a, None, None, None, 0, 2, 0)])  # 11.09 20:00, 12.09 08:00
        self.assertEqual(per_mtu, [SummaryRow(None, self.mtu_a, None, None, None, 0, 2, 0)])

    def test_stock_at_range_end_is_the_stock_rule(self):
        opened = self.station("Янги", self.mtu_a, 9, opening_remaining=7, opening_at=tk(12, 8))
        later = self.station("Кейинги", self.mtu_a, 10, opening_remaining=4, opening_at=tk(13, 8))
        scope = Station.objects.filter(pk__in=[self.a1.pk, opened.pk, later.pk]).order_by("sort")
        rows, (mtu_row,) = summary.summary(scope, D_FROM, D_TO, now=LATER)
        self.assertEqual([r.remaining_end for r in rows], [6, 7, None])
        self.assertEqual(
            {r.station.id: r.remaining_end for r in rows}, reports.remaining_at(scope, periods.ends_for(D_TO)[1])
        )
        self.assertEqual(mtu_row.remaining_end, 13)

    def test_daily_series_has_one_point_per_day(self):
        self.assertEqual(
            summary.daily_series(self.scope(), D_FROM, D_TO, now=LATER),
            {
                "labels": ["10.09.2026", "11.09.2026", "12.09.2026"],
                "unloaded": [3, 2, 9],
                "arrived": [1, 5, 7],
                "remaining": [9, 10, 4],  # stock at 20:00; unknown stations add nothing
                "missing": [3, 4, 4],
            },
        )

    def test_daily_series_remaining_agrees_with_remaining_at(self):
        series = summary.daily_series(self.scope(), D_FROM - timedelta(days=2), D_TO + timedelta(days=1), now=LATER)
        for i, label in enumerate(series["labels"]):
            day = D_FROM - timedelta(days=2 - i)
            with self.subTest(day=label):
                known = [
                    v for v in reports.remaining_at(self.scope(), periods.ends_for(day)[1]).values() if v is not None
                ]
                self.assertEqual(series["remaining"][i], sum(known) if known else None)

    def test_daily_series_missing_adds_up_to_the_summary(self):
        rows, _ = summary.summary(self.scope(), D_FROM, D_TO, now=LATER)
        series = summary.daily_series(self.scope(), D_FROM, D_TO, now=LATER)
        self.assertEqual(sum(series["missing"]), sum(r.missing for r in rows))

    def test_a_day_without_entries_is_a_gap_not_a_zero(self):
        self.assertEqual(
            summary.daily_series(self.scope(), date(2026, 9, 8), date(2026, 9, 8), now=LATER),
            {"labels": ["08.09.2026"], "unloaded": [None], "arrived": [None], "remaining": [None], "missing": [2]},
        )

    def test_a_day_with_no_ended_period_is_a_gap_not_a_zero(self):
        self.assertEqual(
            summary.daily_series(self.scope(), D_TO, D_TO + timedelta(days=1), now=tk(13, 7)),
            {
                "labels": ["12.09.2026", "13.09.2026"],
                "unloaded": [9, None],
                "arrived": [7, None],
                "remaining": [4, None],
                "missing": [4, None],
            },
        )

    def test_compare_by_mtu_and_by_station(self):
        self.assertEqual(
            summary.compare(self.scope(), D_FROM, D_TO, by="mtu"),
            {"labels": ["Тошкент", "Қўқон"], "unloaded": [9, 5], "arrived": [6, 7]},
        )
        self.assertEqual(
            summary.compare(self.scope(), D_FROM, D_TO, by="station"),
            {"labels": ["Олмазор", "Сергели", "Бета"], "unloaded": [9, None, 5], "arrived": [6, None, 7]},
        )
        self.assertEqual(
            summary.compare(Station.objects.filter(mtu=self.mtu_b), D_FROM, D_TO, by="mtu"),
            {"labels": ["Қўқон"], "unloaded": [5], "arrived": [7]},
        )
        with self.assertRaises(ValueError):
            summary.compare(self.scope(), D_FROM, D_TO, by="day")

    def test_query_count_does_not_grow_with_the_number_of_stations(self):
        def add_stations(count):
            for i in range(Station.objects.count(), Station.objects.count() + count):
                s = self.station(f"Станция {i}", (self.mtu_a, self.mtu_b)[i % 2], 100 + i)
                self.entry(s, tk(9, 20), calc=4)
                self.entry(s, tk(11, 8), unloaded=1, arrived=2, calc=5, override=6)

        def read_everything():
            with self.assertNumQueries(3):
                rows, per_mtu = summary.summary(self.scope(), D_FROM, D_TO, now=LATER)
                names = [(r.station.name, r.mtu.name) for r in rows] + [r.mtu.name for r in per_mtu]
            with self.assertNumQueries(3):
                series = summary.daily_series(self.scope(), D_FROM, D_TO, now=LATER)
            with self.assertNumQueries(2):
                bars = summary.compare(self.scope(), D_FROM, D_TO, by="station")
            return len(rows), len(names), len(series["labels"]), len(bars["labels"])

        add_stations(2)
        self.assertEqual(read_everything(), (5, 7, 3, 5))
        add_stations(18)
        self.assertEqual(read_everything(), (23, 25, 3, 23))


DAY = date(2026, 9, 12)
NIGHT_END, DAY_END = periods.ends_for(DAY)
TEMPLATE = openpyxl.load_workbook(export.TEMPLATE).active


def row_values(sheet, row):
    return [sheet.cell(row, col).value for col in range(1, 10)]


def all_text(sheet):
    return " ".join(str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value is not None)


class DayFixture(Base):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.s1 = cls.station("Олмазор", cls.mtu_a, 1)
        cls.entry(cls.s1, NIGHT_END, unloaded=3, arrived=1, calc=8, override=10)
        cls.entry(cls.s1, DAY_END, unloaded=4, arrived=2, calc=8)
        cls.s2 = cls.station("=1+1", cls.mtu_a, 2)  # the ORM skips validators: a hostile name can still exist
        cls.entry(cls.s2, NIGHT_END, unloaded=5, arrived=0, calc=-5)
        cls.s4 = cls.station("Сергели", cls.mtu_a, 3)
        cls.entry(cls.s4, tk(11, 20), calc=6)
        cls.entry(cls.s4, DAY_END, unloaded=1, arrived=1, calc=6)
        cls.s3 = cls.station("Бета", cls.mtu_b, 1)
        cls.boss = User.objects.create_user("boshqarma", password="x" * 16)
        cls.tashkent = User.objects.create_user("toshkent", password="x" * 16, mtu=cls.mtu_a)

    def login(self, user):
        self.client.force_login(user)
        session = self.client.session
        session["staff_login_at"] = timezone.now().timestamp()
        session.save()


class ExportTests(DayFixture):
    def sheet(self, stations, user):
        data = export.export_daily(stations, DAY, user=user, now=LATER)
        return openpyxl.load_workbook(BytesIO(data)).active

    def test_header_is_the_template_with_the_date_in_the_title(self):
        sheet = self.sheet(self.scope(), self.boss)
        self.assertEqual(sheet["A1"].value, f"{TEMPLATE['A1'].value} 12.09.2026")
        for row in (2, 3):
            self.assertEqual(row_values(sheet, row), row_values(TEMPLATE, row))
        self.assertLessEqual(
            {str(r) for r in TEMPLATE.merged_cells.ranges}, {str(r) for r in sheet.merged_cells.ranges}
        )
        self.assertEqual(sheet.column_dimensions["B"].width, TEMPLATE.column_dimensions["B"].width)

    def test_boshqarma_file_is_grouped_by_mtu_with_subtotals_and_a_total(self):
        sheet = self.sheet(self.scope(), self.boss)
        self.assertEqual(sheet["A4"].value, "Тошкент МТУ")
        self.assertIn("A4:I4", sheet.merged_cells)
        self.assertEqual(row_values(sheet, 5), [1, "Олмазор", 3, 1, 10, 4, 2, 8, 7])
        self.assertEqual(row_values(sheet, 6), [2, "=1+1", 5, 0, -5, None, None, -5, "5*"])
        self.assertEqual(row_values(sheet, 7), [3, "Сергели", None, None, 6, 1, 1, 6, "1*"])
        self.assertEqual(row_values(sheet, 8), [None, "Тошкент МТУ жами", 8, 1, 11, 5, 3, 9, "13*"])
        self.assertEqual(sheet["A9"].value, "Қўқон МТУ")
        self.assertEqual(row_values(sheet, 10), [1, "Бета", None, None, None, None, None, None, None])
        self.assertEqual(row_values(sheet, 11), [None, "Қўқон МТУ жами", 0, 0, 0, 0, 0, 0, "0*"])
        self.assertEqual(row_values(sheet, 12), [None, "Жами", 8, 1, 11, 5, 3, 9, "13*"])

    def test_cells_carry_numbers_text_and_marks(self):
        sheet = self.sheet(self.scope(), self.boss)
        self.assertEqual((sheet["C5"].data_type, sheet["I5"].data_type), ("n", "n"))
        name = sheet["B6"]
        self.assertEqual((name.value, name.data_type), ("=1+1", "s"))  # text, never a formula
        self.assertEqual((sheet["I6"].value, sheet["I6"].data_type), ("5*", "s"))

        override = sheet["E5"]
        self.assertEqual(override.comment.text, "Ҳисобланган: 8")
        self.assertEqual(override.fill.fgColor.rgb, export.YELLOW)
        self.assertIsNone(sheet["H5"].comment)

        self.assertEqual(sheet["E6"].font.color.rgb, export.RED)  # negative
        self.assertEqual(sheet["H6"].font.color.rgb, export.RED)  # negative and carried: red wins
        self.assertEqual(sheet["E7"].font.color.rgb, export.GREY)  # carried
        self.assertEqual(sheet["H7"].font.color, TEMPLATE["H5"].font.color)  # entered here: the template's ink
        self.assertIsNone(sheet["C7"].value)  # missing is blank, never 0
        self.assertTrue(sheet["B8"].font.b)
        self.assertEqual(sheet["C5"].border.left.style, TEMPLATE["C5"].border.left.style)
        self.assertEqual(sheet["C5"].alignment.horizontal, TEMPLATE["C3"].alignment.horizontal)

    def test_footer_states_time_account_missing_stations_and_the_star(self):
        text = all_text(self.sheet(self.scope(), self.boss))
        self.assertIn("Яратилган: 01.10.2026 08:00", text)
        self.assertIn("Аккаунт: boshqarma", text)
        self.assertIn("Киритмаган станциялар сони: 3", text)  # =1+1, Сергели, Бета
        self.assertIn("12* — бир давр киритилмаган", text)

    def test_mtu_file_has_its_stations_then_the_total(self):
        sheet = self.sheet(access.stations_for(self.tashkent), self.tashkent)
        self.assertEqual([sheet.cell(r, 2).value for r in range(4, 8)], ["Олмазор", "=1+1", "Сергели", "Жами"])
        self.assertEqual(row_values(sheet, 7), [None, "Жами", 8, 1, 11, 5, 3, 9, "13*"])
        self.assertNotIn("Бета", all_text(sheet))
        self.assertNotIn("Қўқон", all_text(sheet))
        self.assertIn("Киритмаган станциялар сони: 2", all_text(sheet))

    def test_a_period_not_ended_is_blank(self):
        sheet = openpyxl.load_workbook(
            BytesIO(export.export_daily(self.scope(), DAY, user=self.boss, now=tk(12, 10)))
        ).active
        self.assertEqual(row_values(sheet, 5), [1, "Олмазор", 3, 1, 10, None, None, None, 3])
        # The totals of the running day period are blank too, as on the page: nobody entered them yet.
        self.assertEqual(row_values(sheet, 8), [None, "Тошкент МТУ жами", 8, 1, 11, None, None, None, "8*"])
        self.assertEqual(row_values(sheet, 12), [None, "Жами", 8, 1, 11, None, None, None, "8*"])

    def test_a_control_character_in_a_stored_name_does_not_break_the_file(self):
        # The forms refuse such a name now; one stored earlier, or through the ORM, must not make every export fail.
        self.station("Альфа\x0bбекат", self.mtu_a, 4)
        text = all_text(self.sheet(access.stations_for(self.tashkent), self.tashkent))
        self.assertIn("Альфабекат", text)


class ViewTests(DayFixture):
    def get(self, name, **params):
        return self.client.get(reverse(name), params)

    def workbook_text(self, response):
        return all_text(openpyxl.load_workbook(BytesIO(response.content)).active)

    def test_anonymous_is_sent_to_staff_login(self):
        for name in ("staff_report", "staff_export"):
            with self.subTest(name=name):
                self.assertRedirects(self.get(name), reverse("staff_login"), fetch_redirect_response=False)

    def test_export_of_an_mtu_user_has_only_its_stations(self):
        self.login(self.tashkent)
        response = self.get("staff_export", d="2026-09-12", mtu=self.mtu_b.id)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], export.CONTENT_TYPE)
        self.assertEqual(response["Content-Disposition"], 'attachment; filename="don-2026-09-12.xlsx"')
        text = self.workbook_text(response)
        self.assertIn("Олмазор", text)
        self.assertNotIn("Бета", text)

    def test_export_defaults_to_the_latest_ended_period_and_honours_mtu_for_boshqarma(self):
        self.login(self.boss)
        today = periods.label(periods.last_ended(timezone.now()))[0]
        for d in (None, "abc", "2026-02-30", "1900-01-01"):
            with self.subTest(d=d):
                response = self.get("staff_export", **({"d": d} if d else {}))
                self.assertEqual(response["Content-Disposition"], f'attachment; filename="don-{today:%Y-%m-%d}.xlsx"')
        text = self.workbook_text(self.get("staff_export", d="2026-09-12", mtu=self.mtu_b.id))
        self.assertIn("Бета", text)
        self.assertNotIn("Олмазор", text)
        # A day that has not begun is clamped, as the daily page does, not exported as an empty «complete» file.
        response = self.get("staff_export", d="2099-01-01")
        self.assertEqual(response["Content-Disposition"], f'attachment; filename="don-{today:%Y-%m-%d}.xlsx"')

    def test_report_page_for_boshqarma_compares_mtus(self):
        self.login(self.boss)
        response = self.get("staff_report", **{"from": "2026-09-10", "to": "2026-09-12"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["chart"]["compare"]["labels"], ["Тошкент", "Қўқон"])
        self.assertEqual(response.context["chart"]["daily"]["labels"], ["10.09.2026", "11.09.2026", "12.09.2026"])
        self.assertEqual([r.station for r in response.context["station_rows"]], [self.s1, self.s2, self.s4, self.s3])
        self.assertEqual([r.mtu for r in response.context["mtu_rows"]], [self.mtu_a, self.mtu_b])
        self.assertContains(response, 'id="chart-data"')
        self.assertContains(response, "hisobot/chart.umd.min.js")
        self.assertContains(response, "hisobot/charts.js")
        self.assertContains(response, '<th scope="row">Бета</th>', html=True)
        self.assertNotContains(response, "None")  # Бета entered nothing: dashes, not zeros or None
        # 6 periods per station: Олмазор 4, =1+1 5, Сергели 4, Бета 6 not entered
        self.assertContains(response, "Оралиқда киритилмаган станция-даврлар: 19")
        self.assertContains(response, f'href="{reverse("staff_export")}?d=2026-09-12"')
        self.assertContains(response, "Excel юклаб олиш")

    def test_report_page_for_an_mtu_user_compares_its_stations(self):
        self.login(self.tashkent)
        response = self.get("staff_report", **{"from": "2026-09-10", "to": "2026-09-12", "mtu": self.mtu_b.id})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["chart"]["compare"]["labels"], ["Олмазор", "=1+1", "Сергели"])
        self.assertNotContains(response, "Бета")
        self.assertContains(response, "Оралиқда киритилмаган станция-даврлар: 13")

    def test_boshqarma_with_one_mtu_compares_its_stations(self):
        self.login(self.boss)
        response = self.get("staff_report", **{"from": "2026-09-12", "to": "2026-09-12", "mtu": self.mtu_b.id})
        self.assertEqual(response.context["chart"]["compare"]["labels"], ["Бета"])
        self.assertContains(response, f'href="{reverse("staff_export")}?d=2026-09-12&amp;mtu={self.mtu_b.id}"')

    def test_presets_and_defaults(self):
        self.login(self.tashkent)
        last = periods.label(periods.last_ended(timezone.now()))[0]
        for params, d_from in [
            ({}, last - timedelta(days=6)),
            ({"from": "abc", "to": "2026-13-01"}, last - timedelta(days=6)),
        ]:
            with self.subTest(params=params):
                response = self.get("staff_report", **params)
                self.assertEqual((response.context["d_from"], response.context["d_to"]), (d_from, last))
                self.assertEqual([p["current"] for p in response.context["presets"]], [True, False])
        week, month = self.get("staff_report").context["presets"]
        self.assertEqual((week["name"], month["name"]), ("Ҳафта", "Ой"))
        self.assertIn(f"from={last - timedelta(days=29):%Y-%m-%d}", month["url"])

    def test_reversed_or_too_long_range_is_clamped_with_a_notice(self):
        self.login(self.tashkent)
        response = self.get("staff_report", **{"from": "2026-09-12", "to": "2026-09-10"})
        self.assertEqual((response.context["d_from"], response.context["d_to"]), (D_FROM, D_TO))
        self.assertContains(response, 'class="notice-warn"')

        response = self.get("staff_report", **{"from": "2024-01-01", "to": "2026-09-12"})
        self.assertEqual((response.context["d_from"], response.context["d_to"]), (date(2025, 9, 12), D_TO))
        self.assertEqual(len(response.context["chart"]["daily"]["labels"]), 366)
        self.assertContains(response, "366")
        self.assertContains(response, 'class="notice-warn"')

        response = self.get("staff_report", **{"from": "2025-09-12", "to": "2026-09-12"})
        self.assertEqual(response.context["d_from"], date(2025, 9, 12))
        self.assertNotContains(response, 'class="notice-warn"')

    def test_a_future_end_is_clamped_to_the_latest_report_date(self):
        self.login(self.tashkent)
        last = periods.label(periods.last_ended(timezone.now()))[0]
        response = self.get("staff_report", **{"from": "2099-12-25", "to": "2099-12-31"})
        self.assertEqual((response.context["d_from"], response.context["d_to"]), (last, last))
        self.assertContains(response, 'class="notice-warn"')
        self.assertContains(response, f'max="{last:%Y-%m-%d}"', count=2)
        response = self.get("staff_report", to="2099-12-31")
        self.assertEqual((response.context["d_from"], response.context["d_to"]), (last - timedelta(days=6), last))

    def test_grey_means_expected_and_not_entered(self):
        later = self.station("Кейин", self.mtu_a, 9, active_from=tk(20, 8))  # expects nothing in the range
        self.login(self.boss)
        response = self.get("staff_report", **{"from": "2026-09-10", "to": "2026-09-12"})
        html = response.content.decode()
        row = html[html.index(f"<th scope=\"row\">{later.name}</th>"):]
        row = row[: row.index("</tr>")]
        self.assertNotIn("is-missing", row)
        self.assertIn("0 / 0", row)
        row = html[html.index('<th scope="row">Бета</th>'):]
        self.assertIn("is-missing", row[: row.index("</tr>")])  # expected 6, entered none
